"""Match R365 "Missing Vendor Item" placeholder invoice lines to counted items.

R365 books a vendor line it cannot match to an item as a placeholder
("Missing Vendor Item - Wine Cost (5320)") carrying only qty and price, so the
variance cannot credit the purchase to anything the bar counts. h.wood AP's
daily "Items Missing in Archimedes" email lists the vendor's item codes for
exactly those lines (one code per placeholder line). This script:

  1. reads every AP report .xlsx in --reports (save the email attachments there),
  2. pulls the venue's placeholder lines that are NOT yet in kount_invoice_line_map,
  3. resolves each code through KevaOS vendor_items / other venues' invoice lines,
     ONLY within the invoice's own vendor family (an item code means nothing
     across distributors),
  4. pairs each line with a code on the SAME invoice by unit price (case or bottle),
  5. for lines left over, accepts a price-only match only when exactly one counted
     item fits (cost x pack within 2%) and the vendor sells wine/spirits,
  6. never matches a different bottle size than the one the venue counts,
  7. writes a review workbook and an insert script. Nothing is written to the
     database by this script -- review the workbook, then run the .sql.

Same rules as the 2026-10-08 load (migrations 0071/0072, 77 Alphabet rows).

Usage (from anywhere; needs the supabase CLI linked in --repo):
  python match_placeholders.py --reports "C:/path/to/AP reports" \
      --repo "C:/Github Projects/Restaurant-App" --out "C:/path/to/output"
Defaults are for Alphabet: --kount-venue v12 --ops-venue 4d2f6062-c696-49f0-9356-ac4c0c8c7b0d
"""
import argparse, json, math, pathlib, re, subprocess, sys, tempfile, zipfile, difflib, datetime
import xml.etree.ElementTree as ET

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
PACKS = (1, 2, 3, 6, 12, 24)
FAMILY = [('southern', 'souther', 'sgws'), ('harbor',), ('breakthru',), ('anheuser',), ('pacific edge',),
          ('beaune',), ('chambers',), ('martine',), ('duckhorn',), ('republic national', 'rndc'), ('johnson brothers',)]
ALCOHOL = ('southern', 'souther', 'beaune', 'chambers', 'martine', 'mascot', 'duckhorn', 'pacific', 'breakthru',
           'harbor', 'parker', 'skurnik', 'beauchamp', 'republic national', 'johnson brothers', 'el gato')
SIZE_RE = re.compile(r'(\d*\.?\d+)\s*(ml|l|lt|oz|fl\.?\s?oz|gal)\b', re.I)


# ── helpers ───────────────────────────────────────────────────────────────────
def family(name):
    n = (name or '').lower()
    return next((f[0] for f in FAMILY if any(k in n for k in f)), n)


def size_ml(s):
    s = (s or '').lower()
    m = re.search(r'size:\s*(\d*\.?\d+)\s*(ml|l)\b', s) or SIZE_RE.search(s)
    if not m:
        return None
    n = float(m.group(1)); u = m.group(2).replace(' ', '').replace('.', '')
    f = {'ml': 1, 'l': 1000, 'lt': 1000, 'oz': 29.5735, 'floz': 29.5735, 'gal': 3785.41}.get(u)
    return round(n * f) if f else None


def tokens(s):
    s = re.sub(r'\(bpc:[^)]*\)', ' ', (s or '').lower())
    s = SIZE_RE.sub(' ', s)
    s = re.sub(r"[’'`\",()/\-.&+*]", ' ', s)
    s = re.sub(r'\b\d+(\.\d+)?\b', ' ', s)
    return {t for t in s.split() if len(t) > 1}


def similar(a, b):
    ta, tb = tokens(a), tokens(b)
    if not ta or not tb:
        return 0.0
    j = len(ta & tb) / len(ta | tb)
    r = difflib.SequenceMatcher(None, ' '.join(sorted(ta)), ' '.join(sorted(tb))).ratio()
    return max(j, r)


def fit(price, target):
    return abs(price - target) / price if price > 0 and target > 0 else 9.0


def q(s):
    return "'" + str(s).replace("'", "''") + "'"


def db(repo, sql):
    """Run read-only SQL through the linked supabase CLI and return rows."""
    with tempfile.NamedTemporaryFile('w', suffix='.sql', delete=False, encoding='utf-8') as f:
        f.write(sql)
        path = f.name
    out = subprocess.run(['supabase', 'db', 'query', '--linked', '-f', path], cwd=repo,
                         capture_output=True, shell=sys.platform == 'win32')
    raw = out.stdout.decode('utf-8', 'replace')
    if '{' not in raw:
        sys.exit('query failed: ' + out.stderr.decode('utf-8', 'replace')[-600:])
    return json.loads(raw[raw.index('{'):raw.rindex('}') + 1])['rows']


# ── 1. AP reports ─────────────────────────────────────────────────────────────
def read_xlsx_rows(path):
    z = zipfile.ZipFile(path)
    ss = []
    if 'xl/sharedStrings.xml' in z.namelist():
        ss = [''.join(t.text or '' for t in si.iter('{%s}t' % NS['m']))
              for si in ET.fromstring(z.read('xl/sharedStrings.xml')).findall('m:si', NS)]
    wb = ET.fromstring(z.read('xl/workbook.xml'))
    rels = {r.get('Id'): r.get('Target') for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
    for s in wb.find('m:sheets', NS):
        tgt = rels[s.get('{%s}id' % NS['r'])]
        p = tgt.lstrip('/') if tgt.startswith(('/xl', 'xl/')) else 'xl/' + tgt
        for row in ET.fromstring(z.read(p)).iter('{%s}row' % NS['m']):
            cells = []
            for c in row.findall('m:c', NS):
                v = c.find('m:v', NS); t = c.get('t')
                val = ss[int(v.text)] if (t == 's' and v is not None) else (
                    ''.join(x.text or '' for x in c.iter('{%s}t' % NS['m'])) if t == 'inlineStr' else (v.text if v is not None else ''))
                cells.append(val or '')
            yield cells


def parse_reports(folder):
    """-> {document#: set(codes)} from every 'Items Missing' report in the folder."""
    codes = {}
    for f in sorted(pathlib.Path(folder).glob('*.xlsx')):
        hdr = None
        for cells in read_xlsx_rows(f):
            low = [c.strip().lower() for c in cells]
            if 'document#' in low and 'exceptioncomment' in low:
                hdr = {n: i for i, n in enumerate(low)}
                continue
            if not hdr or len(cells) <= hdr['exceptioncomment']:
                continue
            doc = cells[hdr['document#']].strip()
            comment = cells[hdr['exceptioncomment']]
            # codes are the tokens after each ':' -- digits-bearing, 3+ chars
            for seg in comment.split(':')[1:]:
                seg = seg.split('Item')[0]
                for tok in re.split(r'[,\s;]+', seg):
                    tok = tok.strip().strip("'").strip('.')
                    if len(tok) >= 3 and re.search(r'\d', tok) and not re.match(r'^\(?\d{4}\)?$', tok):
                        codes.setdefault(doc, set()).add(tok)
    return codes


# ── 2-6. matching ─────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--reports', required=True)
    ap.add_argument('--repo', default='C:/Github Projects/Restaurant-App')
    ap.add_argument('--out', default='.')
    ap.add_argument('--kount-venue', default='v12')
    ap.add_argument('--ops-venue', default='4d2f6062-c696-49f0-9356-ac4c0c8c7b0d')
    a = ap.parse_args()
    out = pathlib.Path(a.out); out.mkdir(parents=True, exist_ok=True)
    stamp = datetime.date.today().isoformat()

    RC = parse_reports(a.reports)
    print(f'AP reports: {len(RC)} documents with codes')

    lines = db(a.repo, f"""
select il.id::text line_id, i.invoice_number inv, i.invoice_date::text d, v.name vendor, il.gl_code gl,
       il.qty::text qty, il.unit_cost::text unit_cost, il.line_total::text line_total
  from invoice_lines il join invoices i on i.id = il.invoice_id left join vendors v on v.id = i.vendor_id
 where i.venue_id = '{a.ops_venue}' and i.external_source = 'r365_odata' and il.gl_code ~ '^53'
   and il.description ilike 'missing vendor item%' and il.line_total > 0
   and not exists (select 1 from kount_invoice_line_map lm where lm.venue_id = '{a.kount_venue}' and lm.invoice_line_id = il.id)
 order by i.invoice_date;""")
    print(f'unmatched placeholder lines: {len(lines)}')
    if not lines:
        return

    # One code = one placeholder line. Codes already consumed by a line loaded
    # earlier (evidence 'code NNN ...') must not be offered again on that invoice,
    # or the same purchase is credited twice.
    used = db(a.repo, f"""
select i.invoice_number inv, substring(lm.evidence from '^code ([^ ]+)') code
  from kount_invoice_line_map lm join invoice_lines il on il.id = lm.invoice_line_id join invoices i on i.id = il.invoice_id
 where lm.venue_id = '{a.kount_venue}' and lm.evidence ~ '^code ';""")
    n_used = 0
    for u in used:
        if u['code'] and u['inv'] in RC and u['code'] in RC[u['inv']]:
            RC[u['inv']].discard(u['code']); n_used += 1
    print(f'codes already used by loaded lines (excluded): {n_used}')

    counted = db(a.repo, f"""
with c as (select distinct e.master_item_id m from kount_entries e join kount_audits au on au.id = e.audit_id
            where au.venue_id = '{a.kount_venue}' and au.status <> 'cancelled' and e.master_item_id is not null)
select mi.id::text mid, mi.name,
       coalesce((select cost_per_unit from kount_venue_cost_overrides o where o.venue_id = '{a.kount_venue}' and o.master_item_id = mi.id),
                (select avg_cost from purchase_items p where p.master_item_id = mi.id and p.avg_cost > 0 order by updated_at desc nulls last limit 1))::text cost
  from c join master_items mi on mi.id = c.m where mi.is_active;""")
    cost = {c['name']: float(c['cost']) for c in counted if c['cost']}
    mid = {c['name']: c['mid'] for c in counted}

    allcodes = sorted({c for v in RC.values() for c in v})
    res = []
    if allcodes:
        vals = ','.join(f"({q(c)})" for c in allcodes)
        res = db(a.repo, f"""
with c(code) as (values {vals})
select vi.vendor_item_code code, v.name vendor, mi.name master, vi.last_price::text last_price, vi.units_per_pack::text upp
  from vendor_items vi join c on c.code = vi.vendor_item_code join vendors v on v.id = vi.vendor_id
  left join purchase_items pi on pi.id = vi.item_id left join master_items mi on mi.id = pi.master_item_id where mi.name is not null
union all
(select distinct on (il.vendor_item_code, v.name) il.vendor_item_code, v.name, mi.name, il.unit_cost::text, null
  from invoice_lines il join c on c.code = il.vendor_item_code join invoices i on i.id = il.invoice_id join vendors v on v.id = i.vendor_id
  left join purchase_items pi on pi.id = il.item_id left join master_items mi on mi.id = coalesce(il.master_item_id, pi.master_item_id)
 where mi.name is not null order by il.vendor_item_code, v.name, i.invoice_date desc);""")
    code_info = {}
    for r in res:
        code_info.setdefault(r['code'], []).append(r)

    def counted_name(master):
        if master in cost:
            return master
        best = max(cost, key=lambda n: similar(master, n))
        if similar(master, best) < 0.85:
            return None
        sa, sb = size_ml(master), size_ml(best)
        return None if (sa and sb and abs(sa - sb) > 5) else best

    def pack_for(price, name, hint=None):
        c = cost.get(name)
        if not c:
            return (hint or 1), 9.0
        best = min(PACKS, key=lambda p: fit(price, p * c))
        if hint and fit(price, hint * c) <= fit(price, best * c) + 0.05:
            best = hint
        return best, fit(price, best * c)

    result = []
    by_inv = {}
    for l in lines:
        by_inv.setdefault(l['inv'], []).append(l)
    for inv, ls in by_inv.items():
        vfam = family(ls[0]['vendor'])
        codes = RC.get(inv, set())
        pairs = []
        for l in ls:
            U = float(l['unit_cost'])
            for code in codes:
                for r in code_info.get(code, []):
                    if family(r['vendor']) != vfam:
                        continue
                    name = counted_name(r['master'])
                    if not name:
                        continue
                    lp = float(r['last_price'] or 0)
                    upp = int(float(r['upp'])) if r['upp'] else None
                    pk, f = pack_for(U, name, upp if upp and lp and fit(U, lp) < 0.15 else None)
                    pairs.append((min(f, fit(U, lp) if lp else 9.0), l['line_id'], code, name, pk))
        pairs.sort()
        used_l, used_c = set(), set()
        for f, lid, code, name, pk in pairs:
            if lid in used_l or code in used_c or f > 0.15:
                continue
            used_l.add(lid); used_c.add(code)
            l = next(x for x in ls if x['line_id'] == lid)
            result.append(dict(l, master=name, pack=pk, units=round(float(l['qty']) * pk, 2), conf='code', how=f'code {code} (price fit {f:.0%})'))
        for l in ls:
            if l['line_id'] in used_l:
                continue
            U = float(l['unit_cost'])
            alc = any(k in (l['vendor'] or '').lower() for k in ALCOHOL)
            cands = {(n, p) for n, c in (cost.items() if alc else []) for p in PACKS if fit(U, p * c) <= 0.02}
            if len(cands) == 1:
                (n, p), = cands
                result.append(dict(l, master=n, pack=p, units=round(float(l['qty']) * p, 2), conf='price', how='price only (unique)'))
            else:
                result.append(dict(l, master=None, pack=None, units=None, conf=None,
                                   how=('ambiguous: ' + '; '.join(f'{n} x{p}' for n, p in sorted(cands))[:200]) if cands
                                   else ('no AP codes for invoice' if not codes else 'unresolved')))

    matched = [r for r in result if r['master']]
    print(f'matched {len(matched)} of {len(result)} (${sum(float(r["line_total"]) for r in matched):,.0f} of ${sum(float(r["line_total"]) for r in result):,.0f})')

    # review workbook
    try:
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = 'Matched'
        hdr = ['Date', 'Vendor', 'Invoice', 'GL', 'Qty', 'Unit cost', 'Line total', 'Counted item', 'Pack', 'Units', 'How', 'OK? (Y / correct item)']
        for title, rows in (('Matched', matched), ('Needs review', [r for r in result if not r['master']])):
            sh = ws if title == 'Matched' else wb.create_sheet(title)
            sh.append(hdr)
            for r in rows:
                sh.append([r['d'], r['vendor'], r['inv'], r['gl'], float(r['qty']), float(r['unit_cost']), float(r['line_total']),
                           r['master'], r['pack'], r['units'], r['how'], ''])
        wb.save(out / f'placeholder-matches-{a.kount_venue}-{stamp}.xlsx')
    except ImportError:
        json.dump(result, open(out / f'placeholder-matches-{a.kount_venue}-{stamp}.json', 'w', encoding='utf-8'), indent=1)

    # insert script (matched only); review first, then run with supabase db query --linked -f
    tag = f'placeholder-matcher {stamp}'
    vals = ',\n'.join(f"  ({q(r['line_id'])}::uuid, {q(mid[r['master']])}::uuid, {r['units']}, {q(r['conf'])}, {q(r['how'])})" for r in matched)
    sql = f"""-- Placeholder matches for {a.kount_venue}, {stamp}: {len(matched)} lines. Review the workbook before running.
-- Undo: delete from public.kount_invoice_line_map where venue_id = '{a.kount_venue}' and set_by = {q(tag)};
begin;
insert into public.kount_invoice_line_map (venue_id, invoice_line_id, master_item_id, units, confidence, evidence, set_by)
select '{a.kount_venue}', s.line_id, s.mid, s.units, s.conf, s.ev, {q(tag)}
  from (values
{vals}
  ) s(line_id, mid, units, conf, ev)
on conflict do nothing;
commit;
"""
    if matched:
        (out / f'placeholder-matches-{a.kount_venue}-{stamp}.sql').write_text(sql, encoding='utf-8', newline='\n')
    print('wrote', out)


if __name__ == '__main__':
    main()
